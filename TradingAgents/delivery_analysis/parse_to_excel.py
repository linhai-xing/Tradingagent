import json, re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import datetime

RAW = r"C:\Users\72955\Desktop\tradingagent\TradingAgents\ocr_raw.json"
OUT = r"C:\Users\72955\Desktop\tradingagent\TradingAgents\delivery_orders.xlsx"

with open(RAW, encoding="utf-8") as f:
    all_data = json.load(f)

trades = []
seen = set()  # dedup key

for filename, items in all_data.items():
    texts = [item["text"] for item in items]

    # Group into trade entries
    # Each trade: [trade_type+name, price, amount, date_time, qty, fee]
    i = 0
    while i < len(texts):
        # Skip header/menu items
        if texts[i] in ("同花顺App", "对账单", "价格/数量", "金额/税费①", "Y",
                        "本月", "近三月", "近半年", "今年", "全部", "自定义",
                        "银行转证券", "证券转银行"):
            # For bank transfers, also skip the following numeric fields
            if texts[i] in ("银行转证券", "证券转银行"):
                i += 1
                # Skip: 0.000, amount, 银-date, 0, 0.00
                while i < len(texts) and re.match(r'^[\d.,\-]+$', texts[i].replace(",", "")) or texts[i].startswith("银"):
                    i += 1
                continue
            i += 1
            continue

        # Skip numeric noise (time, battery %, app prefix)
        if re.match(r'^\d{2}:\d{2}$', texts[i]):  # time like 09:10
            i += 1
            continue
        if texts[i] in ("!!!", ":!!!", "111", "63", "0", "60:60"):
            i += 1
            continue
        if re.match(r'^\d+$', texts[i]) and len(texts[i]) <= 2:  # single/double digit
            i += 1
            continue

        # Check if this is a trade start
        if texts[i].startswith("证券买入-") or texts[i].startswith("证券卖出-"):
            if i + 5 < len(texts):
                trade_type_full = texts[i]
                price_str = texts[i+1]
                amount_str = texts[i+2]
                date_str = texts[i+3]
                qty_str = texts[i+4]
                fee_str = texts[i+5]

                # Validate: price should be a decimal
                if not re.match(r'^[\d,.]+$', price_str.replace(",", "").replace(".", "")):
                    i += 1
                    continue

                # Parse direction and stock name
                if trade_type_full.startswith("证券买入-"):
                    direction = "买入"
                    stock_name = trade_type_full[5:]
                else:
                    direction = "卖出"
                    stock_name = trade_type_full[5:]

                # Parse date: 买05-2609:44 or 卖05-2609:44
                date_match = re.match(r'[买卖银](\d{2})-(\d{2})(\d{2}:\d{2})', date_str)
                if date_match:
                    month = date_match.group(1)
                    day = date_match.group(2)
                    time = date_match.group(3)
                    trade_date = f"2026-{month}-{day} {time}"
                else:
                    trade_date = date_str

                # Parse numbers
                try:
                    price = float(price_str.replace(",", ""))
                except ValueError:
                    price = None

                try:
                    amount = float(amount_str.replace(",", ""))
                except ValueError:
                    amount = None

                try:
                    qty = int(float(qty_str.replace(",", "")))
                except ValueError:
                    qty = None

                try:
                    fee = float(fee_str.replace(",", ""))
                except ValueError:
                    fee = None

                # Create dedup key
                dedup_key = f"{trade_date}|{direction}|{stock_name}|{price}|{qty}|{amount}"
                if dedup_key not in seen:
                    seen.add(dedup_key)
                    trades.append({
                        "日期": trade_date,
                        "方向": direction,
                        "股票名称": stock_name,
                        "价格": price,
                        "数量(股)": qty,
                        "成交金额": amount,
                        "税费": fee,
                    })

                i += 6
                continue

        # Check for monthly summary line (e.g., "2026-05")
        if re.match(r'^2026-\d{2}$', texts[i]):
            i += 1
            continue

        # Skip percentage/summary lines
        if re.match(r'^[-+]\d', texts[i]) and '%' not in texts[i]:
            i += 1
            continue
        if '%' in texts[i] or '上证' in texts[i]:
            i += 1
            continue

        i += 1

# Sort by date descending
trades.sort(key=lambda t: t["日期"], reverse=True)

# Create Excel
wb = Workbook()
ws = wb.active
ws.title = "交割单"

# Headers
headers = ["序号", "日期", "方向", "股票名称", "价格(元)", "数量(股)", "成交金额(元)", "税费(元)"]
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)

for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border

# Column widths
col_widths = [6, 18, 8, 14, 12, 10, 16, 12]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

# Data rows
buy_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
sell_fill = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
data_font = Font(name="微软雅黑", size=10)

for idx, t in enumerate(trades, 1):
    row = idx + 1
    values = [
        idx,
        t["日期"],
        t["方向"],
        t["股票名称"],
        t["价格"],
        t["数量(股)"],
        t["成交金额"],
        t["税费"],
    ]
    row_fill = buy_fill if t["方向"] == "买入" else sell_fill
    for col, val in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=val)
        cell.font = data_font
        cell.border = thin_border
        cell.fill = row_fill
        if col in (1, 3):
            cell.alignment = Alignment(horizontal="center")
        elif col >= 5:
            cell.alignment = Alignment(horizontal="right")
            if isinstance(val, float):
                cell.number_format = '#,##0.00'

# Summary row
summary_row = len(trades) + 3
ws.cell(row=summary_row, column=1, value="汇总").font = Font(name="微软雅黑", bold=True, size=11)

total_buy = sum(t["成交金额"] for t in trades if t["方向"] == "买入" and t["成交金额"])
total_sell = sum(t["成交金额"] for t in trades if t["方向"] == "卖出" and t["成交金额"])
total_fee = sum(t["税费"] for t in trades if t["税费"])

ws.cell(row=summary_row, column=2, value=f"总成交 {len(trades)} 笔").font = Font(name="微软雅黑", bold=True, size=11)
ws.cell(row=summary_row+1, column=2, value="买入总额").font = Font(name="微软雅黑", size=10)
ws.cell(row=summary_row+1, column=7, value=abs(total_buy)).number_format = '#,##0.00'
ws.cell(row=summary_row+2, column=2, value="卖出总额").font = Font(name="微软雅黑", size=10)
ws.cell(row=summary_row+2, column=7, value=total_sell).number_format = '#,##0.00'
ws.cell(row=summary_row+3, column=2, value="税费合计").font = Font(name="微软雅黑", size=10)
ws.cell(row=summary_row+3, column=8, value=total_fee).number_format = '#,##0.00'
ws.cell(row=summary_row+4, column=2, value="净盈亏").font = Font(name="微软雅黑", bold=True, size=11)
net = total_sell + total_buy  # total_buy is negative
ws.cell(row=summary_row+4, column=7, value=net).number_format = '#,##0.00'
ws.cell(row=summary_row+4, column=7).font = Font(name="微软雅黑", bold=True, size=11,
    color="FF0000" if net >= 0 else "008000")

ws.freeze_panes = "A2"

wb.save(OUT)
print(f"Excel saved: {OUT}")
print(f"Total trades: {len(trades)}")
print(f"Buy: {sum(1 for t in trades if t['方向']=='买入')}, "
      f"Sell: {sum(1 for t in trades if t['方向']=='卖出')}")
print(f"Net P&L: {net:,.2f}")

# Print summary
print("\n--- Trade Summary ---")
for t in trades:
    print(f"{t['日期'][:10]} {t['方向']:2s} {t['股票名称']:6s} "
          f"@{t['价格']:>10.3f} x {t['数量(股)']:>5d} = {t['成交金额']:>12.2f}")
